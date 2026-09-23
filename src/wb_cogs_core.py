#!/usr/bin/env python3
"""
Разбор и загрузка файла себестоимости по неделям ("СС <проект> от <дата>.xlsx")
в ClickHouse — см. schema_wb_cogs.sql.

Формат файла (лист "CC общ"): широкая матрица, артикулы в строках, недели в
столбцах. Три строки шапки: номер недели, дата начала недели, дата конца
недели (в той же ячейке слева — подпись "Артикул"). Дальше каждая строка —
артикул и себестоимость единицы товара в каждую из недель.

Загрузка идемпотентна: ReplacingMergeTree по (sku, week_start), повторная
заливка того же или более свежего файла заменяет строки, а не задваивает их.
Артикул, пропавший из новой выгрузки, при этом останется от старой — это
осознанно (история себестоимости не должна пропадать из-за того, что товар
вывели из ассортимента), но означает, что "удалить артикул" через повторную
загрузку файла нельзя, только вручную.
"""

import logging
import os
from datetime import date, timedelta
from pathlib import Path

import clickhouse_connect
import openpyxl

logger = logging.getLogger(__name__)

DEFAULT_SHEET = "CC общ"
COLUMNS = ["sku", "week_start", "week_end", "week_label", "unit_cost", "sku_source", "source_file"]


def get_client():
    host = os.environ["CLICKHOUSE_HOST"]
    port = int(os.environ.get("CLICKHOUSE_PORT", "8443"))
    user = os.environ.get("CLICKHOUSE_USER", "default")
    password = os.environ["CLICKHOUSE_PASSWORD"]
    database = os.environ.get("CLICKHOUSE_DATABASE", "default")
    secure = os.environ.get("CLICKHOUSE_SECURE", "1") != "0"
    return clickhouse_connect.get_client(
        host=host, port=port, username=user, password=password,
        database=database, secure=secure,
    )


def _as_date(value) -> date:
    if isinstance(value, date):
        return value if not hasattr(value, "date") else value.date()
    raise ValueError(f"ожидалась дата, получено {value!r}")


def parse_weeks(header_nums, header_begin, header_end) -> list:
    """Столбцы-недели из трёх строк шапки. Возвращает (col_idx, label, begin, end)."""
    weeks = []
    for col in range(1, len(header_begin)):
        begin_raw = header_begin[col] if col < len(header_begin) else None
        end_raw = header_end[col] if col < len(header_end) else None
        if begin_raw is None and end_raw is None:
            continue  # хвостовой пустой столбец — в файле от 14.09.26 такой есть
        if begin_raw is None or end_raw is None:
            raise ValueError(f"столбец {col + 1}: заполнена только одна из дат недели "
                             f"(начало={begin_raw!r}, конец={end_raw!r})")
        begin, end = _as_date(begin_raw), _as_date(end_raw)
        if begin.weekday() != 0:
            raise ValueError(f"неделя {begin}..{end} начинается не с понедельника")
        if (end - begin).days != 6:
            raise ValueError(f"неделя {begin}..{end} длиной {(end - begin).days + 1} дн., ожидалось 7")
        label_raw = header_nums[col] if col < len(header_nums) else None
        label = "" if label_raw is None else str(int(label_raw) if isinstance(label_raw, float) else label_raw)
        weeks.append((col, label, begin, end))

    if not weeks:
        raise ValueError("в шапке не найдено ни одной недели — проверьте лист и формат файла")

    for prev, cur in zip(weeks, weeks[1:]):
        if cur[2] - prev[3] != timedelta(days=1):
            raise ValueError(f"разрыв между неделями: {prev[2]}..{prev[3]} и {cur[2]}..{cur[3]}")
    return weeks


def parse_file(path, sheet_name: str = DEFAULT_SHEET) -> list:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"файл не найден: {path}")

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        raise RuntimeError(f"не удалось открыть {path.name}: {e}") from e

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"в файле нет листа {sheet_name!r}, есть: {wb.sheetnames}")
        rows = list(wb[sheet_name].iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < 4:
        raise ValueError(f"лист {sheet_name!r}: меньше 4 строк, нет ни шапки, ни данных")

    weeks = parse_weeks(rows[0], rows[1], rows[2])
    logger.info("Недель в файле: %d (%s..%s)", len(weeks), weeks[0][2], weeks[-1][3])

    out = []
    skus_seen = set()
    skipped_cells = 0
    for row in rows[3:]:
        raw_sku = row[0] if row else None
        if raw_sku is None or not str(raw_sku).strip():
            continue
        sku_source = str(raw_sku).strip()
        sku = sku_source.lower()
        if sku in skus_seen:
            logger.warning("Артикул %r встречается в файле повторно — строки будут схлопнуты в ReplacingMergeTree", sku_source)
        skus_seen.add(sku)

        for col, label, begin, end in weeks:
            value = row[col] if col < len(row) else None
            if value is None or value == "":
                skipped_cells += 1
                continue
            try:
                unit_cost = float(value)
            except (TypeError, ValueError):
                logger.warning("Артикул %r, неделя %s: значение %r не число — строка пропущена",
                               sku_source, begin, value)
                skipped_cells += 1
                continue
            out.append([sku, begin, end, label, unit_cost, sku_source, path.name])

    logger.info("Артикулов: %d, строк к загрузке: %d, пустых/некорректных ячеек пропущено: %d",
                len(skus_seen), len(out), skipped_cells)
    if not out:
        raise ValueError("в файле не нашлось ни одной строки себестоимости")
    return out


def ingest_file(path, sheet_name: str = DEFAULT_SHEET, client=None) -> dict:
    rows = parse_file(path, sheet_name)
    client = client or get_client()
    try:
        client.insert("wb_cogs_weekly", rows, column_names=COLUMNS)
    except Exception as e:
        raise RuntimeError(f"не удалось записать себестоимость в ClickHouse: {e}") from e

    weeks = {r[1] for r in rows}
    skus = {r[0] for r in rows}
    return {"rows": len(rows), "skus": len(skus), "weeks": len(weeks),
            "week_min": min(weeks), "week_max": max(weeks)}
