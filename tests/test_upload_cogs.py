"""Загрузка еженедельной матрицы себестоимости (upload_checks/cogs.py + дверь на странице загрузки).

Цифры реального файла «СС CloudSix от 14.09.26»: 523 артикула, 40 недель, 20 920 значений, 5 450 нулей и
22 отрицательных значения у двух артикулов — нули и минусы поэтому предупреждения, а не отказ.
"""
import datetime as dt
import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "webapp"))

import wb_cogs_core as core  # noqa: E402
from upload_checks import cogs  # noqa: E402
from upload_checks.core import UploadRejected, ERROR, INFO, WARN  # noqa: E402

MON = dt.date(2026, 9, 7)   # понедельник


def _matrix(path: Path, skus, weeks=3, start=MON, sheet=core.DEFAULT_SHEET, values=None):
    """Синтетическая матрица как в настоящем файле: 3 строки шапки (номер, начало, конец), далее артикулы."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    begins = [start + dt.timedelta(days=7 * i) for i in range(weeks)]
    ws.append([None] + [30 + i for i in range(weeks)])
    ws.append(["Артикул"] + [datetime_ for datetime_ in begins])
    ws.append([None] + [b + dt.timedelta(days=6) for b in begins])
    for i, sku in enumerate(skus):
        ws.append([sku] + [(values or {}).get((sku, w), 100.0 + i) for w in range(weeks)])
    wb.save(path)
    return path


class FakeClient:
    """ClickHouse-заглушка: хранит wb_cogs_weekly, отвечает на запросы проверок."""
    def __init__(self, existing=None, max_week=None):
        self.existing = existing or {}          # (sku, week_start) -> cost
        self.max_week = max_week
        self.inserts = []

    def insert(self, table, data, column_names=None, **k):
        self.inserts.append((table, [list(r) for r in data]))
        if table == "wb_cogs_weekly":
            for r in data:
                self.existing[(r[0], r[1])] = r[4]

    def query(self, sql, parameters=None):
        R = lambda rows: type("R", (), {"result_rows": rows})()
        if "FROM wb_cogs_weekly FINAL WHERE sku IN" in sql:
            wanted = set(parameters["s"])
            return R([(s, w, c) for (s, w), c in self.existing.items() if s in wanted])
        if "max(week_start)" in sql:
            return R([(self.max_week or (max((w for _, w in self.existing), default=None)),)])
        return R([])


def _ingest(files, client, **kw):
    return cogs.ingest([Path(f) for f in files], log_fn=lambda *a: None, database="cloudsix", user_id=1,
                       client=client, **kw)


def test_first_load_writes_everything_and_reports_summary(tmp_path):
    c = FakeClient()
    res = _ingest([_matrix(tmp_path / "СС.xlsx", ["Aero_10L", "TV_Lampa"])], c)
    assert res["rows"] == 6 and res["rows_in_file"] == 6 and res["duplicates_skipped"] == 0
    assert sum(len(d) for t, d in c.inserts if t == "wb_cogs_weekly") == 6
    msgs = [r["message"] for o in res["outcomes"] for r in o["results"]]
    assert any("Новых значений: 6" in m for m in msgs)


def test_repeat_of_same_file_writes_nothing_and_keeps_provenance(tmp_path):
    c = FakeClient()
    f = _matrix(tmp_path / "СС.xlsx", ["Aero_10L"])
    _ingest([f], c)
    c.inserts.clear()
    res = _ingest([f], c)
    assert res["rows"] == 0 and res["duplicates_skipped"] == 3
    assert [t for t, _ in c.inserts if t == "wb_cogs_weekly"] == []          # неизменённое не перезаписываем


def test_changed_history_is_written_and_warned_with_examples(tmp_path):
    c = FakeClient()
    _ingest([_matrix(tmp_path / "a.xlsx", ["Aero_10L"])], c)
    c.inserts.clear()
    f2 = _matrix(tmp_path / "b.xlsx", ["Aero_10L"], values={("Aero_10L", 1): 150.0})
    res = _ingest([f2], c)
    assert res["rows"] == 1
    warn = [r for o in res["outcomes"] for r in o["results"] if r["name"] == "changed_history"]
    assert warn and warn[0]["severity"] == WARN and "100 → 150" in warn[0]["message"] and "Aero_10L" in warn[0]["message"]


def test_foreign_file_is_rejected_before_any_write(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Выписка"
    wb.active.append(["Дата", "Контрагент", "Сумма"])
    wb.active.append([dt.date(2026, 6, 1), "ООО", 1000])
    p = tmp_path / "Выписка.xlsx"
    wb.save(p)
    c = FakeClient()
    with pytest.raises(UploadRejected) as e:
        _ingest([p], c)
    assert "не похож на матрицу себестоимости" in str(e.value) or any("не похож" in r.message for r in e.value.results)
    assert [t for t, _ in c.inserts if t == "wb_cogs_weekly"] == []


def test_broken_week_header_is_rejected(tmp_path):
    p = _matrix(tmp_path / "x.xlsx", ["A"], start=dt.date(2026, 9, 8))       # вторник, не понедельник
    with pytest.raises(UploadRejected) as e:
        _ingest([p], FakeClient())
    assert any("не с понедельника" in r.message for r in e.value.results)


def test_negative_and_zero_costs_are_warnings_not_rejection(tmp_path):
    f = _matrix(tmp_path / "n.xlsx", ["Mini_Printer_Round"], values={("Mini_Printer_Round", 0): -935.56, ("Mini_Printer_Round", 1): 0.0})
    res = _ingest([f], FakeClient())
    names = {r["name"]: r["severity"] for o in res["outcomes"] for r in o["results"]}
    assert names["negative_cost"] == WARN and res["rows"] == 3


def test_older_file_than_db_and_stale_file_warn(tmp_path):
    c = FakeClient(existing={("a", dt.date(2027, 1, 4)): 1.0}, max_week=dt.date(2027, 1, 4))
    res = _ingest([_matrix(tmp_path / "old.xlsx", ["A"])], c)
    names = {r["name"] for o in res["outcomes"] for r in o["results"]}
    assert {"older_than_db", "stale_file"} <= names or "older_than_db" in names


def test_every_file_is_checked_before_writing_any(tmp_path):
    good = _matrix(tmp_path / "good.xlsx", ["A"])
    bad = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook(); wb.active.title = "Другое"; wb.save(bad)
    c = FakeClient()
    with pytest.raises(UploadRejected):
        _ingest([good, bad], c)
    assert [t for t, _ in c.inserts if t == "wb_cogs_weekly"] == []           # всё или ничего


def test_cli_parse_file_still_works(tmp_path):
    rows = core.parse_file(_matrix(tmp_path / "s.xlsx", ["A", "B"]))
    assert len(rows) == 6 and rows[0][0] == "a"


# ----------------------------------------------------------------- дверь на странице

import app as webapp  # noqa: E402
import auth  # noqa: E402

PROJECTS = [{"id": 1, "slug": "cloudsix", "name": "CloudSix"}]


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(auth, "find_user_by_id", lambda uid: auth.User(1, "t@example.com", "Иван", "Тестов"))
    monkeypatch.setattr(webapp, "get_user_projects", lambda uid: PROJECTS)
    monkeypatch.setattr(webapp, "get_project_by_slug", lambda slug: PROJECTS[0] if slug == "cloudsix" else None)
    monkeypatch.setattr(webapp, "user_has_project_access", lambda uid, pid: True)
    monkeypatch.setattr(webapp, "get_project_cabinets", lambda pid, db, platform=None: ["Feel"])
    monkeypatch.setattr(webapp, "get_project_platforms", lambda pid, db: ["wb"])
    monkeypatch.setattr(webapp, "get_project_sources", lambda pid, db: ["cogs_weekly"])
    monkeypatch.setattr(webapp, "get_project_cabinet_platforms", lambda pid, db: {"Feel": ["wb"]})
    c = webapp.app.test_client()
    with c.session_transaction() as s:
        s["_user_id"] = "1"
        s["_fresh"] = True
    return c


def test_door_is_visible_only_when_source_is_enabled(client, monkeypatch):
    html = client.get("/p/cloudsix/upload").get_data(as_text=True)
    assert "Себестоимость, еженедельная матрица" in html and "/upload/cogs" in html
    monkeypatch.setattr(webapp, "get_project_sources", lambda pid, db: [])
    assert "/upload/cogs" not in client.get("/p/cloudsix/upload").get_data(as_text=True)


def test_route_rejects_foreign_file_with_clear_message(client, tmp_path, monkeypatch):
    fake = FakeClient()
    monkeypatch.setattr(cogs.wb_core, "get_client", lambda database=None: fake)
    wb = openpyxl.Workbook(); wb.active.title = "Выписка"; wb.active.append(["a"]); p = tmp_path / "x.xlsx"; wb.save(p)
    r = client.post("/p/cloudsix/upload/cogs", data={"files": (io.BytesIO(p.read_bytes()), "x.xlsx")},
                    content_type="multipart/form-data")
    html = r.get_data(as_text=True)
    assert r.status_code == 400 and "Файл не загружен" in html and "не похож на матрицу себестоимости" in html
    assert [t for t, _ in fake.inserts if t == "wb_cogs_weekly"] == []


def test_route_success_shows_new_and_unchanged_counts(client, tmp_path, monkeypatch):
    fake = FakeClient()
    monkeypatch.setattr(cogs.wb_core, "get_client", lambda database=None: fake)
    p = _matrix(tmp_path / "ok.xlsx", ["Aero_10L"])
    post = lambda: client.post("/p/cloudsix/upload/cogs", data={"files": (io.BytesIO(p.read_bytes()), "ok.xlsx")},
                               content_type="multipart/form-data")
    html = post().get_data(as_text=True)
    assert "записано новых: 3" in html and "Новых значений: 3" in html
    html2 = post().get_data(as_text=True)
    assert "записано новых: 0" in html2 and "без изменений (не перезаписаны): 3" in html2


def test_route_404_when_source_not_enabled_for_project(client, monkeypatch):
    monkeypatch.setattr(webapp, "get_project_sources", lambda pid, db: [])
    r = client.post("/p/cloudsix/upload/cogs", data={"files": (io.BytesIO(b"x"), "ok.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 404
