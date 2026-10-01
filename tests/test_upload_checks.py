"""Проверки и дедуп при ручной загрузке (ТЗ 04, блок A): ядро, дедуп, конвейер WB/Ozon."""

import datetime as dt
from collections import Counter
from pathlib import Path

import openpyxl
import pytest

import ozon_core
import wb_core
from upload_checks import core, dedup, ozon as ozon_source, wb as wb_source
from upload_checks.core import UploadRejected, ERROR, INFO, WARN


# ---------------------------------------------------------------- вспомогательное

class FakeClient:
    """Минимальный ClickHouse: хранит вставки, на SELECT ... FROM <table> FINAL
    отдаёт сохранённые строки нужных колонок (WHERE игнорируется)."""

    def __init__(self):
        self.tables = {}          # table -> (column_names, rows)
        self.insert_calls = []

    def insert(self, table, data, column_names=None):
        self.insert_calls.append(table)
        cols, rows = self.tables.setdefault(table, (list(column_names), []))
        for row in data:
            rows.append(dict(zip(column_names, row)))

    def query(self, sql, parameters=None):
        import re
        m = re.match(r"SELECT (.+) FROM (\w+) FINAL", sql)
        cols = [c.strip() for c in m.group(1).split(",")]
        stored = self.tables.get(m.group(2), ([], []))[1]
        rows = [tuple(r.get(c) for c in cols) for r in stored]
        return type("R", (), {"result_rows": rows})()


def _write_xlsx(path: Path, core_module, header_row: int, drop=(), rename=None, rows=3):
    """Синтетический xlsx: заголовки = первые алиасы маппинга, значения по типу колонки."""
    rename = rename or {}
    with open(core_module.MAPPING_PATH, encoding="utf-8") as f:
        import yaml
        cols = yaml.safe_load(f)["columns"]
    headers, canon_order = [], []
    for canon, info in cols.items():
        if canon in drop:
            continue
        headers.append(rename.get(canon, info["aliases"][0]))
        canon_order.append((canon, info["type"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(header_row):
        ws.append(["Период: тест"])
    ws.append(headers)
    for i in range(1, rows + 1):
        vals = []
        for canon, t in canon_order:
            if canon == "row_num":
                vals.append(i)
            elif t == "Float64":
                vals.append(10.5 * i)
            elif t == "Int32":
                vals.append(i)
            elif t == "Date":
                vals.append(dt.date(2026, 9, i))
            else:
                vals.append(f"{canon}-{i}")
        ws.append(vals)
    wb.save(path)


@pytest.fixture
def client(monkeypatch):
    c = FakeClient()
    monkeypatch.setattr(wb_core, "get_client", lambda database=None: c)
    monkeypatch.setattr(ozon_core, "get_client", lambda database=None: c)
    import reconcile_wb
    monkeypatch.setattr(reconcile_wb, "run_reconciliation", lambda *a, **k: [])
    return c


# ------------------------------------------------------------------------ dedup

def test_fingerprint_same_for_file_row_and_db_row():
    file_row = {"a": "x", "b": 10.5, "c": dt.date(2026, 9, 1), "d": None, "e": 3}
    db_row = ("x", 10.5, dt.date(2026, 9, 1), None, 3)
    assert dedup.row_fingerprint(file_row, list("abcde")) == dedup.fingerprint(db_row)


def test_split_new_full_partial_and_intra_file_duplicates():
    cols = ["k"]
    existing = Counter({dedup.fingerprint(["a"]): 1})
    rows = [{"k": "a"}, {"k": "b"}, {"k": "b"}]       # a — дубль, две одинаковых b — законные
    new, dups = dedup.split_new(rows, existing, cols)
    assert dups == 1 and [r["k"] for r in new] == ["b", "b"]


def test_split_new_second_file_sees_first():
    cols = ["k"]
    existing = Counter()
    dedup.split_new([{"k": "a"}], existing, cols)
    new, dups = dedup.split_new([{"k": "a"}], existing, cols)
    assert new == [] and dups == 1


# ------------------------------------------------------------------------- core

def _specs_aliases(core_module):
    a2c, _ = core_module.load_mapping()
    return a2c, core.load_column_specs(core_module.MAPPING_PATH)


def test_missing_money_column_is_error_missing_optional_is_not():
    a2c, specs = _specs_aliases(wb_core)
    headers = [h for canon, s in specs.items() for h in
               [next(k for k, v in a2c.items() if v == canon)]
               if canon not in ("payable_to_seller", "chrt_id")]
    res = core.check_headers(headers, a2c, specs)
    err = [r for r in res if r.severity == ERROR]
    assert len(err) == 1 and "payable_to_seller" in err[0].details["columns"]
    assert all("chrt_id" not in r.details.get("columns", []) for r in res)


def test_unmapped_header_is_warning():
    a2c, specs = _specs_aliases(wb_core)
    headers = [next(k for k, v in a2c.items() if v == c) for c in specs] + ["Новая колонка"]
    res = core.check_headers(headers, a2c, specs)
    assert [r.severity for r in res] == [WARN] and "Новая колонка" in res[0].message


def test_check_cabinet_only_from_project_list():
    assert core.check_cabinet("Feel", ["Feel", "ARB"]) == []
    res = core.check_cabinet("Fel", ["Feel", "ARB"])
    assert res[0].severity == ERROR and "Fel" in res[0].message


def test_persist_swallows_journal_failure():
    class Broken:
        def insert(self, *a, **k):
            raise RuntimeError("нет таблицы")
    o = core.FileOutcome("f.xlsx", 1, 1, 0)
    assert core.persist(Broken(), [o], user_id=1, project="p", cabinet="c", source="s") is False


# ------------------------------------------------------------------- конвейер WB

def test_wb_upload_then_reupload_writes_nothing(tmp_path, client):
    f = tmp_path / "Отчёт №111222333_1.xlsx"
    _write_xlsx(f, wb_core, 0)
    first = wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix", user_id=7)
    assert first["rows"] == 3 and first["duplicates_skipped"] == 0

    again = wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix", user_id=7)
    assert again["rows"] == 0 and again["duplicates_skipped"] == 3
    msgs = [r["message"] for r in again["outcomes"][0]["results"]]
    assert any("уже был загружен целиком" in m for m in msgs)
    assert len(client.tables["wb_reports"][1]) == 3          # данные не задвоены
    assert "upload_checks" in client.tables                  # журнал ведётся


def test_wb_partial_overlap_writes_only_new(tmp_path, client):
    f1 = tmp_path / "Отчёт №111_1.xlsx"
    _write_xlsx(f1, wb_core, 0, rows=3)
    wb_source.ingest([f1], "Feel", log_fn=lambda *_: None, database="cloudsix")
    f2 = tmp_path / "Отчёт №111_2.xlsx"                      # тот же отчёт, но уже 5 строк
    _write_xlsx(f2, wb_core, 0, rows=5)
    res = wb_source.ingest([f2], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert res["rows"] == 2 and res["duplicates_skipped"] == 3


def test_wb_renamed_money_column_rejected_and_nothing_written(tmp_path, client):
    f = tmp_path / "Отчёт №111_1.xlsx"
    _write_xlsx(f, wb_core, 0, rename={"payable_to_seller": "Совсем новое название"})
    with pytest.raises(UploadRejected) as e:
        wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert any(r.name == "missing_money_columns" for r in e.value.results)
    assert "wb_reports" not in client.tables                 # ни одной строки данных


def test_wb_one_bad_file_rejects_whole_batch(tmp_path, client):
    good = tmp_path / "Отчёт №1_1.xlsx"
    bad = tmp_path / "без_номера.xlsx"
    _write_xlsx(good, wb_core, 0)
    _write_xlsx(bad, wb_core, 0)
    with pytest.raises(UploadRejected) as e:
        wb_source.ingest([good, bad], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert any(r.name == "report_number_missing" for r in e.value.results)
    assert "wb_reports" not in client.tables


def test_unreadable_file_rejected(tmp_path, client):
    f = tmp_path / "Отчёт №5_1.xlsx"
    f.write_bytes(b"not an xlsx")
    with pytest.raises(UploadRejected) as e:
        wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert any(r.name == "unreadable_file" for r in e.value.results)


def test_wb_summary_mismatch_reported_after_write(tmp_path, client, monkeypatch):
    import reconcile_wb
    rec = [("Feel", 111, "x", None, None, "payable", 1.0, 2.0, 1.0, 0.01, 0)]
    monkeypatch.setattr(reconcile_wb, "run_reconciliation", lambda *a, **k: rec)
    f = tmp_path / "Отчёт №111_1.xlsx"
    _write_xlsx(f, wb_core, 0)
    res = wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert res["rows"] == 3                                    # данные записаны
    r = [x for x in res["outcomes"][0]["results"] if x["name"] == "summary_reconciliation"][0]
    assert r["severity"] == ERROR and "не сходятся" in r["message"]


# ----------------------------------------------------------------- конвейер Ozon

def test_ozon_same_data_under_other_filename_is_duplicate(tmp_path, client):
    a, b = tmp_path / "accr_jan.xlsx", tmp_path / "accr_jan_copy.xlsx"
    _write_xlsx(a, ozon_core, 1)
    _write_xlsx(b, ozon_core, 1)
    assert ozon_source.ingest([a], "X-Tech", log_fn=lambda *_: None, database="cloudsix")["rows"] == 3
    again = ozon_source.ingest([b], "X-Tech", log_fn=lambda *_: None, database="cloudsix")
    assert again["rows"] == 0 and again["duplicates_skipped"] == 3
    assert len(client.tables["ozon_reports"][1]) == 3


def test_ozon_missing_money_column_rejected(tmp_path, client):
    f = tmp_path / "accr.xlsx"
    _write_xlsx(f, ozon_core, 1, drop=("total_amount",))
    with pytest.raises(UploadRejected):
        ozon_source.ingest([f], "X-Tech", log_fn=lambda *_: None, database="cloudsix")
    assert "ozon_reports" not in client.tables


# ------------------------------------------------------------------- вебапп (A)

def _webapp():
    import sys
    ROOT = Path(__file__).resolve().parent.parent
    sys.path.insert(0, str(ROOT / "webapp"))
    import app as webapp
    return webapp


def test_reject_unknown_cabinet_renders_400_with_reason(monkeypatch):
    webapp = _webapp()
    monkeypatch.setattr(webapp, "get_project_cabinets", lambda pid, db, platform=None: ["Feel", "ARB"])
    with webapp.app.test_request_context():
        webapp.g.project = {"id": 1, "slug": "cloudsix"}
        resp, status = webapp.reject_unknown_cabinet("cloudsix", "Fel", "wb")
        body = resp if isinstance(resp, str) else resp.get_data(as_text=True)
    assert status == 400
    assert "Fel" in body and "не относится к этому проекту" in body
    assert "в базу ничего не записано" in body


def test_known_cabinet_passes(monkeypatch):
    webapp = _webapp()
    monkeypatch.setattr(webapp, "get_project_cabinets", lambda pid, db, platform=None: ["Feel"])
    with webapp.app.test_request_context():
        webapp.g.project = {"id": 1, "slug": "cloudsix"}
        assert webapp.reject_unknown_cabinet("cloudsix", "Feel", "wb") is None


def test_cabinet_check_uses_platform(monkeypatch):
    """Кабинет Ozon-платформы не годится для WB-формы и наоборот."""
    webapp = _webapp()
    seen = []
    monkeypatch.setattr(webapp, "get_project_cabinets",
                        lambda pid, db, platform=None: seen.append(platform) or [])
    with webapp.app.test_request_context():
        webapp.g.project = {"id": 1, "slug": "cloudsix"}
        webapp.reject_unknown_cabinet("cloudsix", "X-Tech", "ozon")
    assert seen == ["ozon"]


def test_result_page_shows_duplicates_and_checks():
    webapp = _webapp()
    summary = {"files": 1, "rows": 2, "rows_in_file": 5, "duplicates_skipped": 3, "unmapped_columns": [],
               "outcomes": [{"source_file": "f.xlsx", "rows_in_file": 5, "rows_written": 2,
                             "duplicates_skipped": 3,
                             "results": [{"name": "duplicates_skipped", "severity": "info",
                                          "message": "Пропущено дублей: 3 из 5; записано новых: 2."}]}]}
    with webapp.app.test_request_context():
        body = webapp.render_template("detail_result.html", error=None, summary=summary,
                                      logs=[], slug="cloudsix")
    assert "пропущено дублей: 3" in body and "Пропущено дублей: 3 из 5" in body


# ---------------------------------------------- неактивная колонка «Удержание Агентского НДС»

def test_agent_vat_column_is_optional_float_in_mapping_and_schema():
    specs = core.load_column_specs(wb_core.MAPPING_PATH)
    assert specs["agent_vat_withholding"] == {"type": "Float64", "optional": True}
    a2c, _ = wb_core.load_mapping()
    assert a2c["Удержание Агентского НДС"] == "agent_vat_withholding"
    assert "agent_vat_withholding" in (Path(__file__).resolve().parent.parent / "src" / "schema_wb.sql").read_text(encoding="utf-8")


def test_file_without_agent_vat_column_is_not_rejected_and_with_it_is_not_extra(tmp_path, client):
    f = tmp_path / "Отчёт №555_1.xlsx"
    _write_xlsx(f, wb_core, 0, drop=("agent_vat_withholding",))          # как 377 из 403 реальных файлов
    assert wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix")["rows"] == 3
    g = tmp_path / "Отчёт №556_1.xlsx"
    _write_xlsx(g, wb_core, 0)                                           # с колонкой
    res = wb_source.ingest([g], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert not any(r["name"] == "unmapped_columns" for r in res["outcomes"][0]["results"])
    stored = [r for r in client.tables["wb_reports"][1] if r["report_number"] == 556][0]
    assert stored["agent_vat_withholding"] == 10.5 and "Удержание Агентского НДС" not in stored["extra_columns"]


def test_logistics_coefficient_column_is_optional_float_and_not_extra(tmp_path, client):
    specs = core.load_column_specs(wb_core.MAPPING_PATH)
    assert specs["logistics_coefficient"] == {"type": "Float64", "optional": True}
    f = tmp_path / "Отчёт №777_1.xlsx"
    _write_xlsx(f, wb_core, 0)
    res = wb_source.ingest([f], "Feel", log_fn=lambda *_: None, database="cloudsix")
    assert not any(r["name"] == "unmapped_columns" for r in res["outcomes"][0]["results"])
    stored = client.tables["wb_reports"][1][0]
    assert stored["logistics_coefficient"] == 10.5 and "Коэффициент логистики" not in stored["extra_columns"]
